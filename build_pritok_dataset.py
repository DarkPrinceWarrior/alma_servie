"""
Build pritok (inflow change) dataset from xlsx files.
Produces:
  - db/pritok_anomaly_database_<freq>.csv   (interpolated on uniform grid)
  - db/pritok_intervals.csv                 (ground truth intervals)

Usage:
  python build_pritok_dataset.py                  # default 2min
  python build_pritok_dataset.py --freq 15s
"""
import warnings
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

warnings.filterwarnings('ignore')

WELL_FILES = {
    '3261': 'pritok_change_anomaly_files/3261_Кустовое_Приток.xlsx',
    '495':  'pritok_change_anomaly_files/495_ЮЯ_Приток.xlsx',
    '902':  'pritok_change_anomaly_files/902_ЮЯ_Приток.xlsx',
}

PARAM_RENAME = {
    'Давление на выкиде ЭЦН': 'Давление на приеме насоса кгс/см²',
}


def parse_xlsx(well_id, filepath):
    print(f'  Парсинг {well_id} из {filepath}...')
    wb = openpyxl.load_workbook(filepath, read_only=True)
    series_list = []

    for sname in wb.sheetnames:
        ws = wb[sname]
        rows = list(ws.iter_rows(values_only=True))
        if len(rows) < 3:
            continue

        full_name = str(rows[1][0]) if rows[1][0] else ''
        parts = full_name.rsplit('.', 1)
        param_name = parts[-1].strip() if len(parts) > 1 else full_name.strip()
        param_name = PARAM_RENAME.get(param_name, param_name)

        data_rows = rows[2:]
        times = []
        values = []
        for r in data_rows:
            if r[0] is None:
                continue
            try:
                ts = pd.to_datetime(r[0], dayfirst=True)
                val = float(r[1]) if r[1] is not None else np.nan
                times.append(ts)
                values.append(val)
            except (ValueError, TypeError):
                continue

        if times:
            s = pd.Series(values, index=pd.DatetimeIndex(times), name=param_name)
            s = s[~s.index.duplicated(keep='first')].sort_index()
            series_list.append(s)

    wb.close()
    print(f'    {well_id}: {len(series_list)} параметров из xlsx')
    return series_list


def build_intervals():
    print('Парсинг сводной информации...')
    svod = pd.read_excel('negermet_anomaly_files/Сводная информация.xlsx', header=None)

    header_row = None
    for i in range(min(10, len(svod))):
        vals = [str(v).strip() if pd.notna(v) else '' for v in svod.iloc[i]]
        if '№' in vals and 'Скважина' in vals:
            header_row = i
            break

    if header_row is None:
        raise ValueError('Не найдена строка заголовков в сводной информации')

    data = svod.iloc[header_row + 1:].reset_index(drop=True)

    intervals = []
    for _, row in data.iterrows():
        anom_type = str(row.iloc[3]).strip() if pd.notna(row.iloc[3]) else ''
        if 'приток' not in anom_type.lower():
            continue

        well_id = str(row.iloc[2]).strip().lower()
        data_start = pd.to_datetime(row.iloc[4], errors='coerce')
        data_end = pd.to_datetime(row.iloc[5], errors='coerce')

        s1 = pd.to_datetime(row.iloc[7], errors='coerce')
        e1 = pd.to_datetime(row.iloc[8], errors='coerce')
        if pd.notna(s1) and pd.notna(e1):
            intervals.append({
                'well_id': well_id,
                'start_date': s1,
                'end_date': e1,
                'data_start': data_start,
                'data_end': data_end,
            })

        if len(row) > 10:
            s2 = pd.to_datetime(row.iloc[9], errors='coerce')
            e2 = pd.to_datetime(row.iloc[10], errors='coerce')
            if pd.notna(s2) and pd.notna(e2):
                intervals.append({
                    'well_id': well_id,
                    'start_date': s2,
                    'end_date': e2,
                    'data_start': data_start,
                    'data_end': data_end,
                })

    idf = pd.DataFrame(intervals)
    idf = idf.sort_values(['well_id', 'start_date']).reset_index(drop=True)
    idf['interval_idx'] = idf.groupby('well_id').cumcount() + 1
    return idf


def main(freq='2min'):
    print(f'=== Построение датасета Изменение притока (шаг={freq}) ===\n')

    well_series = {}
    for wid, fpath in WELL_FILES.items():
        slist = parse_xlsx(wid, fpath)
        if slist:
            well_series[wid] = slist

    if not well_series:
        print('Нет данных!')
        return

    print(f'\nПостроение датасета (общий диапазон по параметрам, шаг {freq})...')
    frames = []
    all_param_names = set()

    for wid in sorted(well_series.keys()):
        slist = well_series[wid]

        param_starts = []
        param_ends = []
        for s in slist:
            param_starts.append(s.index.min())
            param_ends.append(s.index.max())

        common_start = max(param_starts)
        common_end = min(param_ends)

        if common_start >= common_end:
            print(f'  {wid}: ПРОПУСК — нет общего диапазона')
            continue

        grid_start = common_start.ceil(freq)
        grid_end = common_end.floor(freq)
        grid = pd.date_range(grid_start, grid_end, freq=freq)

        if len(grid) < 2:
            print(f'  {wid}: ПРОПУСК — слишком короткий общий диапазон')
            continue

        print(f'  {wid}: общий диапазон {common_start} — {common_end} '
              f'({common_end - common_start}), сетка {len(grid)} точек')

        df_well = pd.DataFrame({'timestamp': grid})
        for s in slist:
            param_name = s.name
            all_param_names.add(param_name)
            trimmed = s[(s.index >= common_start) & (s.index <= common_end)]
            if len(trimmed) < 2:
                df_well[param_name] = np.nan
                continue
            combined_idx = trimmed.index.union(grid)
            reindexed = trimmed.reindex(combined_idx).sort_index()
            interpolated = reindexed.interpolate(method='time')
            interpolated = interpolated.ffill().bfill()
            df_well[param_name] = interpolated.reindex(grid).values

        df_well['well_id'] = wid
        frames.append(df_well)

    if not frames:
        print('Нет данных после интерполяции!')
        return

    result = pd.concat(frames, ignore_index=True)
    numeric_cols = sorted(all_param_names)
    result = result[['timestamp', 'well_id'] + [c for c in numeric_cols if c in result.columns]]

    freq_label = freq.replace(' ', '')
    out_path = Path(f'db/pritok_anomaly_database_{freq_label}.csv')
    result.to_csv(out_path, index=False)
    print(f'\nDatabase: {out_path} ({len(result)} rows)')

    idf = build_intervals()
    intervals_path = Path('db/pritok_intervals.csv')
    idf.to_csv(intervals_path, index=False)
    print(f'\nIntervals: {intervals_path}')
    print(idf.to_string())

    print(f'\nВерификация датасета ({freq}):')
    result['timestamp'] = pd.to_datetime(result['timestamp'])
    for wid in sorted(result['well_id'].unique()):
        sub = result[result['well_id'] == wid].sort_values('timestamp')
        deltas = sub['timestamp'].diff().dropna()
        nan_total = sub[numeric_cols].isna().sum().sum() if all(c in sub.columns for c in numeric_cols) else 0
        print(f'  {wid}: median={deltas.median()}, count={len(sub)}, NaN={nan_total}')


if __name__ == '__main__':
    import argparse
    parser = argparse.ArgumentParser()
    parser.add_argument('--freq', default='2min', help='Resample frequency (e.g. 15s, 2min)')
    args = parser.parse_args()
    main(freq=args.freq)
