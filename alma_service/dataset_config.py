from __future__ import annotations

from functools import lru_cache

import openpyxl

from alma_service.paths import MODEL_PARAMS_PATH, NEGREMET_RAW_DIR, PRITOK_RAW_DIR, SALT_RAW_DIR


PARAM_RENAME = {
    "Давление на выкиде ЭЦН": "Давление на приеме насоса кгс/см²",
}

NEGREMET_WELL_FILES = {
    "5271г": NEGREMET_RAW_DIR / "5271г_ЮЯ_ННКТ.xlsx",
    "524": NEGREMET_RAW_DIR / "524_ЮЯ_ННКТ.xlsx",
    "3509г": NEGREMET_RAW_DIR / "3509г_Кустовое_ННКТ.xlsx",
    "1123л": NEGREMET_RAW_DIR / "1123л_Кустовое_ННКТ.xlsx",
    "172г": NEGREMET_RAW_DIR / "172г_Яркое_ННКТ.xlsx",
}

PRITOK_WELL_FILES = {
    "3261": PRITOK_RAW_DIR / "3261_Кустовое_Приток.xlsx",
    "495": PRITOK_RAW_DIR / "495_Южно-Ягунская_Приток.xlsx",
    "902": PRITOK_RAW_DIR / "902_ЮЯ_Приток.xlsx",
    "129л": PRITOK_RAW_DIR / "129л_Приток.xlsx",
}

SALT_WELL_FILES = {
    "3244г": SALT_RAW_DIR / "3244г_Кустовое_Соли.xlsx",
    "3245": SALT_RAW_DIR / "3245_Кустовое_Соли.xlsx",
    "149г": SALT_RAW_DIR / "149г_ВИК_Соли.xlsx",
    "3269": SALT_RAW_DIR / "3269_Кустовое_Соли.xlsx",
    "4039": SALT_RAW_DIR / "4039_Дружное_Соли.xlsx",
    "408": SALT_RAW_DIR / "408_Дружное_Соли.xlsx",
    "3245(2)": SALT_RAW_DIR / "3245(2)_Кустовое_Соли.xlsx",
    "2991г": SALT_RAW_DIR / "2991г_Соли.xlsx",
}

TEST_WELLS = {
    "negermet": {"3509г"},
    "pritok": {"902", "129л"},
    "salt": {"2991г"},
}


@lru_cache(maxsize=1)
def load_model_parameters() -> frozenset[str]:
    wb = openpyxl.load_workbook(MODEL_PARAMS_PATH, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    params = []
    for row in ws.iter_rows(values_only=True):
        value = row[0] if row else None
        if value is None:
            continue
        value = str(value).strip()
        if value:
            params.append(value)
    wb.close()
    return frozenset(params)


def normalize_param_name(name: str) -> str:
    return PARAM_RENAME.get(name, name)


def split_for_well(anomaly_name: str, well_id: str) -> str:
    normalized_well_id = str(well_id).strip().lower()
    return "test" if normalized_well_id in TEST_WELLS.get(anomaly_name, set()) else "train"
