from __future__ import annotations

import itertools
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

from alma_service.anomaly_specs import DatasetSpec
from alma_service.dataset_config import normalize_param_name, split_for_well
from alma_service.paths import DB_DIR, SUMMARY_INFO_PATH, ensure_dir


def _looks_like_datetime(value: object) -> bool:
    if value is None:
        return False
    try:
        parsed = pd.to_datetime(value, dayfirst=True, errors="coerce")
    except Exception:
        return False
    return pd.notna(parsed)


def parse_parameter_series(well_id: str, filepath: Path) -> list[pd.Series]:
    print(f"  Парсинг {well_id} из {filepath}...")
    wb = openpyxl.load_workbook(filepath, read_only=True)
    series_list: list[pd.Series] = []

    for sname in wb.sheetnames:
        ws = wb[sname]
        row_iter = ws.iter_rows(values_only=True)
        _ = next(row_iter, None)
        header_row = next(row_iter, None)
        third_row = next(row_iter, None)
        if header_row is None:
            continue

        full_name = str(header_row[0]) if header_row[0] else ""
        parts = full_name.rsplit(".", 1)
        param_name = parts[-1].strip() if len(parts) > 1 else full_name.strip()
        param_name = normalize_param_name(param_name)

        if _looks_like_datetime(third_row[0] if third_row else None):
            row_iter = itertools.chain([third_row], row_iter)

        times: list[pd.Timestamp] = []
        values: list[float] = []
        for row in row_iter:
            if not row or row[0] is None:
                continue
            try:
                ts = pd.to_datetime(row[0], dayfirst=True)
                val = float(row[1]) if row[1] is not None else np.nan
            except (TypeError, ValueError):
                continue
            times.append(ts)
            values.append(val)

        if not times:
            continue

        series = pd.Series(values, index=pd.DatetimeIndex(times), name=param_name)
        series = series[~series.index.duplicated(keep="first")].sort_index()
        series_list.append(series)

    wb.close()
    print(f"    {well_id}: {len(series_list)} параметров из xlsx")
    return series_list


def build_intervals(spec: DatasetSpec) -> pd.DataFrame:
    print("Парсинг сводной информации...")
    svod = pd.read_excel(SUMMARY_INFO_PATH, header=None)

    header_row = None
    for i in range(min(10, len(svod))):
        vals = [str(v).strip() if pd.notna(v) else "" for v in svod.iloc[i]]
        if "№" in vals and "Скважина" in vals:
            header_row = i
            break

    if header_row is None:
        raise ValueError("Не найдена строка заголовков в сводной информации")

    data = svod.iloc[header_row + 1 :].reset_index(drop=True)

    intervals: list[dict[str, object]] = []
    for _, row in data.iterrows():
        anom_type = str(row.iloc[3]).strip() if pd.notna(row.iloc[3]) else ""
        if spec.summary_match not in anom_type.lower():
            continue

        well_id = str(row.iloc[2]).strip().lower()
        data_start = pd.to_datetime(row.iloc[4], errors="coerce")
        data_end = pd.to_datetime(row.iloc[5], errors="coerce")

        s1 = pd.to_datetime(row.iloc[7], errors="coerce")
        e1 = pd.to_datetime(row.iloc[8], errors="coerce")
        if pd.notna(s1) and pd.notna(e1):
            intervals.append(
                {
                    "well_id": well_id,
                    "start_date": s1,
                    "end_date": e1,
                    "data_start": data_start,
                    "data_end": data_end,
                    "split": split_for_well(spec.anomaly_key, well_id),
                }
            )

        if len(row) > 10:
            s2 = pd.to_datetime(row.iloc[9], errors="coerce")
            e2 = pd.to_datetime(row.iloc[10], errors="coerce")
            if pd.notna(s2) and pd.notna(e2):
                intervals.append(
                    {
                        "well_id": well_id,
                        "start_date": s2,
                        "end_date": e2,
                        "data_start": data_start,
                        "data_end": data_end,
                        "split": split_for_well(spec.anomaly_key, well_id),
                    }
                )

    idf = pd.DataFrame(intervals)
    idf = idf.sort_values(["well_id", "start_date"]).reset_index(drop=True)
    idf["interval_idx"] = idf.groupby("well_id").cumcount() + 1
    return idf


def _build_well_frame(slist: list[pd.Series], freq: str) -> tuple[pd.DataFrame, dict[str, object]] | None:
    starts = [series.index.min() for series in slist]
    ends = [series.index.max() for series in slist]
    union_start = min(starts)
    union_end = max(ends)

    grid_start = union_start.ceil(freq)
    grid_end = union_end.floor(freq)
    grid = pd.date_range(grid_start, grid_end, freq=freq)
    if len(grid) < 2:
        return None

    df_well = pd.DataFrame({"timestamp": grid})
    for series in slist:
        trimmed = series[(series.index >= union_start) & (series.index <= union_end)]
        if len(trimmed) < 2:
            df_well[series.name] = np.nan
            continue
        combined_idx = trimmed.index.union(grid)
        reindexed = trimmed.reindex(combined_idx).sort_index()
        # Strictly causal fill: carry only past information forward to the
        # resampled grid, never interpolate using future sensor values.
        causal_filled = reindexed.ffill()
        df_well[series.name] = causal_filled.reindex(grid).values

    return df_well, {
        "start": union_start,
        "end": union_end,
        "grid_points": len(grid),
    }


def build_dataset(spec: DatasetSpec, freq: str | None = None) -> tuple[pd.DataFrame, pd.DataFrame]:
    freq = freq or spec.default_freq
    ensure_dir(DB_DIR)

    print(f"=== Построение датасета {spec.display_name} (шаг={freq}) ===\n")

    well_series: dict[str, list[pd.Series]] = {}
    for well_id, path in spec.well_files.items():
        slist = parse_parameter_series(well_id, path)
        if slist:
            well_series[well_id] = slist

    if not well_series:
        raise RuntimeError("Нет данных для построения датасета.")

    frames: list[pd.DataFrame] = []
    all_param_names: set[str] = set()
    print(f"\nПостроение датасета (полный диапазон по скважине, шаг {freq})...")
    for well_id in sorted(well_series):
        built = _build_well_frame(well_series[well_id], freq)
        if built is None:
            print(f"  {well_id}: ПРОПУСК — слишком короткий диапазон")
            continue
        df_well, meta = built
        for column in df_well.columns:
            if column not in {"timestamp", "well_id"}:
                all_param_names.add(column)
        df_well["well_id"] = well_id
        frames.append(df_well)
        print(
            f"  {well_id}: диапазон {meta['start']} — {meta['end']} "
            f"({meta['end'] - meta['start']}), сетка {meta['grid_points']} точек"
        )

    if not frames:
        raise RuntimeError("Нет данных после интерполяции.")

    result = pd.concat(frames, ignore_index=True)
    numeric_cols = sorted(all_param_names)
    result = result[["timestamp", "well_id"] + [c for c in numeric_cols if c in result.columns]]

    freq_label = freq.replace(" ", "")
    out_path = DB_DIR / f"{spec.output_prefix}_anomaly_database_{freq_label}.csv"
    result.to_csv(out_path, index=False)
    print(f"\nDatabase: {out_path} ({len(result)} rows)")

    intervals = build_intervals(spec)
    intervals_path = DB_DIR / f"{spec.output_prefix}_intervals.csv"
    intervals.to_csv(intervals_path, index=False)
    print(f"\nIntervals: {intervals_path}")
    print(intervals.to_string())

    print(f"\nВерификация датасета ({freq}):")
    result["timestamp"] = pd.to_datetime(result["timestamp"])
    for well_id in sorted(result["well_id"].unique()):
        sub = result[result["well_id"] == well_id].sort_values("timestamp")
        deltas = sub["timestamp"].diff().dropna()
        usable_cols = [c for c in numeric_cols if not sub[c].isna().all()]
        all_nan_cols = [c for c in numeric_cols if sub[c].isna().all()]
        print(
            f"  {well_id}: median={deltas.median()}, count={len(sub)}, "
            f"usable={len(usable_cols)}, all_nan={len(all_nan_cols)}"
        )
        if all_nan_cols:
            print(f"    {well_id}: all-NaN каналы: {all_nan_cols}")

    return result, intervals
